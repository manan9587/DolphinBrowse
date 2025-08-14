import pandas as pd
import fitz
import docx
from fastapi import FastAPI
from fastapi.responses import FileResponse
from pathlib import Path


class FileProcessor:
    def analyze(self, path: str):
        ext = Path(path).suffix.lower()
        if ext == '.pdf':
            doc = fitz.open(path)
            text = "\n".join(page.get_text() for page in doc)
            return {'text': text}
        if ext == '.docx':
            d = docx.Document(path)
            text = "\n".join(p.text for p in d.paragraphs)
            return {'text': text}
        if ext in ['.xls', '.xlsx']:
            df = pd.read_excel(path, sheet_name=None)
            return {'sheets': {k: df[k].to_dict() for k in df}}
        if ext == '.csv':
            df = pd.read_csv(path)
            return {'data': df.to_dict()}
        raise ValueError('Unsupported format')

    def generate_remarks(self, analysis: dict, out_path: str):
        import xlsxwriter
        wb = xlsxwriter.Workbook(out_path)
        ws = wb.add_worksheet('Remarks')
        ws.write(0, 0, 'Section')
        ws.write(0, 1, 'Content')
        row = 1
        for section, content in analysis.items():
            ws.write(row, 0, section)
            ws.write(row, 1, str(content)[:1000])
            row += 1
        wb.close()
        return out_path


app = FastAPI()


@app.get('/api/files/{file_id}/analyze')
async def analyze_file(file_id: str):
    upload_path = f'/tmp/uploads/{file_id}'
    fp = FileProcessor()
    analysis = fp.analyze(upload_path)
    out = f'/tmp/remarks_{file_id}.xlsx'
    fp.generate_remarks(analysis, out)
    return FileResponse(
        out,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename='remarks.xlsx',
    )
